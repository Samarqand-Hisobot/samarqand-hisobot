import os, io, calendar
from datetime import datetime
from functools import wraps
from flask import Flask, render_template, request, redirect, url_for, session, flash, send_file
from flask_sqlalchemy import SQLAlchemy
from werkzeug.security import check_password_hash, generate_password_hash

app = Flask(__name__)
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY', 'change-this-secret-before-deploying')
database_url = os.environ.get('DATABASE_URL', 'sqlite:///samarqand_hisobot.db')
if database_url.startswith('postgres://'):
    database_url = database_url.replace('postgres://', 'postgresql+psycopg://', 1)
elif database_url.startswith('postgresql://'):
    database_url = database_url.replace('postgresql://', 'postgresql+psycopg://', 1)
app.config['SQLALCHEMY_DATABASE_URI'] = database_url
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

db = SQLAlchemy(app)

MONTHS = ['Yanvar','Fevral','Mart','Aprel','May','Iyun','Iyul','Avgust','Sentabr','Oktabr','Noyabr','Dekabr']

class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(80), unique=True, nullable=False)
    password_hash = db.Column(db.String(255), nullable=False)
    role = db.Column(db.String(20), nullable=False, default='editor')

class Department(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(120), unique=True, nullable=False)
    reports = db.relationship('MonthlyReport', backref='department', cascade='all, delete-orphan')

class MonthlyReport(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    department_id = db.Column(db.Integer, db.ForeignKey('department.id'), nullable=False)
    year = db.Column(db.Integer, nullable=False)
    month = db.Column(db.Integer, nullable=False)
    revenue = db.Column(db.Float, nullable=False, default=0)
    qiyoslov = db.Column(db.Integer, nullable=False, default=0)
    conclusion_plan = db.Column(db.Integer, nullable=False, default=0)
    conclusion_actual = db.Column(db.Integer, nullable=False, default=0)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    __table_args__ = (db.UniqueConstraint('department_id','year','month', name='unique_dept_year_month'),)


def login_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if 'user_id' not in session:
            return redirect(url_for('login'))
        return fn(*args, **kwargs)
    return wrapper


def admin_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if 'user_id' not in session:
            return redirect(url_for('login'))
        user = db.session.get(User, session['user_id'])
        if not user or user.role != 'admin':
            flash('Bu amal faqat administrator uchun.', 'error')
            return redirect(url_for('index'))
        return fn(*args, **kwargs)
    return wrapper


def get_report(dept_id, year, month):
    report = MonthlyReport.query.filter_by(department_id=dept_id, year=year, month=month).first()
    if report is None:
        report = MonthlyReport(department_id=dept_id, year=year, month=month)
        db.session.add(report)
        db.session.commit()
    return report

@app.route('/login', methods=['GET','POST'])
def login():
    if request.method == 'POST':
        user = User.query.filter_by(username=request.form.get('username','').strip()).first()
        if user and check_password_hash(user.password_hash, request.form.get('password','')):
            session.clear(); session['user_id'] = user.id
            return redirect(url_for('index'))
        flash('Login yoki parol noto‘g‘ri.', 'error')
    return render_template('login.html')

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

@app.route('/')
@login_required
def index():
    departments = Department.query.order_by(Department.name).all()
    if not departments:
        d = Department(name='Samarqand bo‘limi'); db.session.add(d); db.session.commit()
        departments = [d]
    dept_id = request.args.get('department', type=int) or departments[0].id
    year = request.args.get('year', 2026, type=int)
    dept = db.session.get(Department, dept_id) or departments[0]
    rows = []
    for month in range(1,13):
        r25 = MonthlyReport.query.filter_by(department_id=dept.id, year=2025, month=month).first()
        r26 = MonthlyReport.query.filter_by(department_id=dept.id, year=2026, month=month).first()
        active = MonthlyReport.query.filter_by(department_id=dept.id, year=year, month=month).first()
        a = r25 if year == 2025 else r26 if year == 2026 else active
        rev25 = r25.revenue if r25 else 0; rev26 = r26.revenue if r26 else 0
        q25 = r25.qiyoslov if r25 else 0; q26 = r26.qiyoslov if r26 else 0
        rows.append({'month':month,'name':MONTHS[month-1], 'revenue':a.revenue if a else 0,
          'qiyoslov':a.qiyoslov if a else 0,'plan':a.conclusion_plan if a else 0,'actual':a.conclusion_actual if a else 0,
          'rev25':rev25,'rev26':rev26,'rev_growth':((rev26-rev25)/rev25*100 if rev25 else None),
          'q25':q25,'q26':q26,'q_growth':((q26-q25)/q25*100 if q25 else None),
          'completion':((a.conclusion_actual/a.conclusion_plan*100) if a and a.conclusion_plan else None)})
    total_revenue = sum(r['revenue'] for r in rows)
    total_q = sum(r['qiyoslov'] for r in rows)
    total_plan = sum(r['plan'] for r in rows)
    total_actual = sum(r['actual'] for r in rows)
    user = db.session.get(User, session['user_id'])
    return render_template('index.html', departments=departments, dept=dept, rows=rows, year=year,
       total_revenue=total_revenue,total_q=total_q,total_plan=total_plan,total_actual=total_actual,
       completion=total_actual/total_plan*100 if total_plan else None,user=user,months=MONTHS)

@app.route('/save', methods=['POST'])
@login_required
def save():
    dept_id = request.form.get('department_id', type=int)
    year = request.form.get('year', type=int)
    if not db.session.get(Department, dept_id) or not year or year < 2000 or year > 2100:
        flash('Bo‘lim yoki yil noto‘g‘ri.', 'error'); return redirect(url_for('index'))
    for month in range(1,13):
        r = get_report(dept_id, year, month)
        for field in ('revenue','qiyoslov','conclusion_plan','conclusion_actual'):
            raw = request.form.get(f'{field}_{month}', '0').replace(' ','').replace(',','.')
            try: val = float(raw or 0)
            except ValueError: val = 0
            if val < 0: val = 0
            setattr(r, field, int(val) if field != 'revenue' else val)
    db.session.commit()
    flash(f'{year}-yil ma’lumotlari saqlandi.', 'success')
    return redirect(url_for('index', department=dept_id, year=year))

@app.route('/departments', methods=['GET','POST'])
@admin_required
def departments():
    if request.method == 'POST':
        name = request.form.get('name','').strip()
        if name and not Department.query.filter_by(name=name).first():
            db.session.add(Department(name=name)); db.session.commit(); flash('Bo‘lim qo‘shildi.','success')
        else: flash('Bo‘lim nomi bo‘sh yoki avvaldan mavjud.','error')
        return redirect(url_for('departments'))
    return render_template('departments.html', departments=Department.query.order_by(Department.name).all())

@app.route('/users', methods=['GET','POST'])
@admin_required
def users():
    if request.method == 'POST':
        username=request.form.get('username','').strip(); password=request.form.get('password',''); role=request.form.get('role','editor')
        if username and len(password)>=8 and role in ('admin','editor') and not User.query.filter_by(username=username).first():
            db.session.add(User(username=username,password_hash=generate_password_hash(password),role=role)); db.session.commit(); flash('Foydalanuvchi yaratildi.','success')
        else: flash('Login takrorlangan yoki parol kamida 8 belgidan iborat bo‘lishi kerak.','error')
        return redirect(url_for('users'))
    return render_template('users.html', users=User.query.order_by(User.username).all())

@app.route('/export/pdf')
@login_required
def export_pdf():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    dept_id=request.args.get('department', type=int); year=request.args.get('year',2026,type=int)
    dept=db.session.get(Department,dept_id) if dept_id else Department.query.first()
    if not dept: flash('Avval bo‘lim yarating.','error'); return redirect(url_for('index'))
    buf=io.BytesIO(); doc=SimpleDocTemplate(buf,pagesize=landscape(A4),rightMargin=24,leftMargin=24,topMargin=24,bottomMargin=24)
    styles=getSampleStyleSheet(); story=[Paragraph(f'{dept.name.upper()} — HISOBOT',styles['Title']),Paragraph(f'2025 va {year}-yil ko‘rsatkichlari',styles['Heading2']),Spacer(1,12)]
    data=[['Oy','2025 tushum','2026 tushum','O‘zgarish %','2025 qiyoslov','2026 qiyoslov','Reja','Amalda','Bajarilish %']]
    for m in range(1,13):
        a=MonthlyReport.query.filter_by(department_id=dept.id,year=2025,month=m).first(); b=MonthlyReport.query.filter_by(department_id=dept.id,year=2026,month=m).first(); c=MonthlyReport.query.filter_by(department_id=dept.id,year=year,month=m).first()
        rev_a=a.revenue if a else 0; rev_b=b.revenue if b else 0; q_a=a.qiyoslov if a else 0; q_b=b.qiyoslov if b else 0
        growth=f'{(rev_b-rev_a)/rev_a*100:.1f}%' if rev_a else '—'; completion=f'{c.conclusion_actual/c.conclusion_plan*100:.1f}%' if c and c.conclusion_plan else '—'
        data.append([MONTHS[m-1],f'{rev_a:,.0f}',f'{rev_b:,.0f}',growth,f'{q_a:,}',f'{q_b:,}',str(c.conclusion_plan if c else 0),str(c.conclusion_actual if c else 0),completion])
    t=Table(data,repeatRows=1,colWidths=[58,78,78,58,65,65,45,48,62]); t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#123B68')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.4,colors.HexColor('#B9C8D8')),('FONTSIZE',(0,0),(-1,-1),7),('ALIGN',(1,1),(-1,-1),'RIGHT'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F0F5FA')]),('BOTTOMPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),6)])); story.append(t)
    doc.build(story); buf.seek(0); return send_file(buf,as_attachment=True,download_name=f'{dept.name}_{year}_hisobot.pdf',mimetype='application/pdf')

@app.route('/export/pptx')
@login_required
def export_pptx():
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    dept_id=request.args.get('department',type=int); year=request.args.get('year',2026,type=int)
    dept=db.session.get(Department,dept_id) if dept_id else Department.query.first()
    if not dept: flash('Avval bo‘lim yarating.','error'); return redirect(url_for('index'))
    prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[0]); slide.shapes.title.text=f'{dept.name} — HISOBOT'; slide.placeholders[1].text=f'2025 va {year}-yil | Rahbariyat uchun'
    slide=prs.slides.add_slide(prs.slide_layouts[5]); slide.shapes.title.text='Oylik ko‘rsatkichlar taqqoslanishi'
    rows=13; cols=6; table=slide.shapes.add_table(rows,cols, Inches(.35), Inches(1.35), Inches(12.5), Inches(5.4)).table
    headers=['Oy','2025 tushum','2026 tushum','2025 qiyoslov','2026 qiyoslov','Xulosa amalda']
    for j,h in enumerate(headers): table.cell(0,j).text=h
    for m in range(1,13):
        a=MonthlyReport.query.filter_by(department_id=dept.id,year=2025,month=m).first(); b=MonthlyReport.query.filter_by(department_id=dept.id,year=2026,month=m).first()
        c=MonthlyReport.query.filter_by(department_id=dept.id,year=year,month=m).first()
        vals=[MONTHS[m-1],f'{(a.revenue if a else 0):,.0f}',f'{(b.revenue if b else 0):,.0f}',str(a.qiyoslov if a else 0),str(b.qiyoslov if b else 0),str(c.conclusion_actual if c else 0)]
        for j,v in enumerate(vals): table.cell(m,j).text=v
    for cell in table.rows[0].cells:
        cell.fill.solid(); cell.fill.fore_color.rgb=RGBColor(18,59,104)
        for p in cell.text_frame.paragraphs:
            for run in p.runs: run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255); run.font.size=Pt(10)
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for run in p.runs: run.font.name='Aptos'
    buf=io.BytesIO(); prs.save(buf); buf.seek(0); return send_file(buf,as_attachment=True,download_name=f'{dept.name}_{year}_hisobot.pptx',mimetype='application/vnd.openxmlformats-officedocument.presentationml.presentation')

with app.app_context():
    db.create_all()
    # Bootstrap an administrator only when environment credentials are supplied.
    admin_user=os.environ.get('ADMIN_USERNAME'); admin_password=os.environ.get('ADMIN_PASSWORD')
    if admin_user and admin_password and not User.query.filter_by(username=admin_user).first():
        db.session.add(User(username=admin_user,password_hash=generate_password_hash(admin_password),role='admin'))
        db.session.commit()
    if not Department.query.first():
        db.session.add(Department(name='Samarqand bo‘limi')); db.session.commit()
    # Load the supplied report's monthly sample data only into a fresh database.
    dept = Department.query.filter_by(name='Samarqand bo‘limi').first()
    if dept and MonthlyReport.query.filter_by(department_id=dept.id).count() == 0:
        rev25 = [30441600,30004800,33364800,39782400,37641065,43335965,38502908,90205444,50320173]
        rev26 = [52460989,45905095,58579141,48538286,33707981,44198171,66223301,48969061,70669306]
        q25 = [401,427,477,501,480,484,410,1146,543]
        q26 = [335,352,325,407,308,345,445,440,509]
        for m in range(1,10):
            db.session.add(MonthlyReport(department_id=dept.id,year=2025,month=m,revenue=rev25[m-1],qiyoslov=q25[m-1],conclusion_plan=0,conclusion_actual=0))
            db.session.add(MonthlyReport(department_id=dept.id,year=2026,month=m,revenue=rev26[m-1],qiyoslov=q26[m-1],conclusion_plan=5 if m >= 8 else 0,conclusion_actual=2 if m == 9 else 0))
        db.session.commit()

if __name__ == '__main__':
    app.run(host='127.0.0.1', port=int(os.environ.get('PORT','5000')), debug=False)
